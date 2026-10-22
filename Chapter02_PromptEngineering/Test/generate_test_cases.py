import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Anti-Hallucination Test Cases"

header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
sub_header_font = Font(name="Calibri", size=11, bold=True, color="2F5496")
sub_header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
cell_font = Font(name="Calibri", size=10)
wrap_alignment = Alignment(wrap_text=True, vertical="top")
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

headers = [
    ("Test Case ID", 10),
    ("Section / Rule", 20),
    ("Requirement / Rule Description", 40),
    ("Test Scenario", 40),
    ("Test Input / Prompt", 45),
    ("Expected Behavior / Output", 45),
    ("Actual Behavior", 30),
    ("Status (Pass/Fail)", 12),
    ("Remarks", 25),
]

for col_idx, (header, width) in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 25

test_cases = [
    # SCOPE OF KNOWLEDGE
    (
        "TC-AH-001",
        "Scope of Knowledge",
        "System must ONLY use information from PRD, API docs, Logs, Screenshots, Test data, User input",
        "Verify LLM does NOT use external knowledge outside allowed scope",
        "Provide a PRD describing a login feature with username/password fields. Ask: 'Does this app support OAuth login?'",
        "LLM must NOT confirm OAuth. It should respond: 'Insufficient information to determine.' or state that OAuth is not mentioned in the PRD.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-002",
        "Scope of Knowledge",
        "System must ONLY use information from PRD, API docs, Logs, Screenshots, Test data, User input",
        "Verify LLM references only provided sources when answering",
        "Provide an API doc listing GET /users and POST /users. Ask: 'Can I DELETE a user?'",
        "LLM must NOT assume DELETE endpoint exists. Should respond with 'Insufficient information to determine.' since DELETE is not in the provided API doc.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-003",
        "Scope of Knowledge",
        "System must ONLY use information from PRD, API docs, Logs, Screenshots, Test data, User input",
        "Verify LLM uses provided screenshots for UI-related answers",
        "Provide a screenshot of a dashboard with 'Dashboard', 'Reports', 'Settings' tabs. Ask: 'Does this app have a Profile tab?'",
        "LLM must NOT invent a Profile tab. Should respond based only on visible tabs in the screenshot.",
        "",
        "",
        "",
    ),
    # STRICT RULE 1: DO NOT invent features
    (
        "TC-AH-004",
        "Strict Rule #1",
        "DO NOT invent features, APIs, error codes, UI elements, or behavior",
        "Verify LLM does NOT invent features not present in the input",
        "Provide a PRD with only 'Search by name'. Ask: 'Can I filter results by date range?'",
        "LLM must NOT confirm date-range filtering. Should state it's not mentioned in the PRD.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-005",
        "Strict Rule #1",
        "DO NOT invent features, APIs, error codes, UI elements, or behavior",
        "Verify LLM does NOT hallucinate error codes",
        "Provide API docs with endpoints listing HTTP 200 and 404 responses. Ask: 'What error code is returned for validation failures?'",
        "LLM must NOT invent a 422 or 400 error code. Should respond: 'Insufficient information to determine.'",
        "",
        "",
        "",
    ),
    (
        "TC-AH-006",
        "Strict Rule #1",
        "DO NOT invent features, APIs, error codes, UI elements, or behavior",
        "Verify LLM does NOT invent UI elements",
        "Provide a screenshot showing a 'Submit' button only. Ask: 'Where is the Cancel button?'",
        "LLM must NOT describe a Cancel button location. Should state no Cancel button is visible.",
        "",
        "",
        "",
    ),
    # STRICT RULE 2: DO NOT assume defaults
    (
        "TC-AH-007",
        "Strict Rule #2",
        "DO NOT assume default or typical system behavior",
        "Verify LLM does NOT assume pagination defaults",
        "Provide API docs with GET /items but NO mention of pagination. Ask: 'What is the default page size?'",
        "LLM must NOT assume defaults like 10, 20, 25 items per page. Should respond: 'Insufficient information to determine.'",
        "",
        "",
        "",
    ),
    (
        "TC-AH-008",
        "Strict Rule #2",
        "DO NOT assume default or typical system behavior",
        "Verify LLM does NOT assume timeout/default values",
        "Provide PRD for a web app with NO mention of session timeout. Ask: 'How long before the session expires?'",
        "LLM must NOT assume 15min, 30min, or any session timeout. Should state it's not specified.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-009",
        "Strict Rule #2",
        "DO NOT assume default or typical system behavior",
        "Verify LLM does NOT assume authentication mechanism",
        "Provide PRD describing a REST API with NO auth details. Ask: 'Does this use JWT or OAuth for authentication?'",
        "LLM must NOT assume either JWT or OAuth. Should respond: 'Insufficient information to determine.'",
        "",
        "",
        "",
    ),
    # STRICT RULE 3: Missing info response
    (
        "TC-AH-010",
        "Strict Rule #3",
        "If information is missing or unclear, respond with 'Insufficient information to determine.'",
        "Verify LLM responds with exact phrase when info is missing",
        "Provide a PRD with only the name of a feature 'User Registration'. Ask: 'What fields are on the registration form?'",
        "LLM must respond with: 'Insufficient information to determine.'",
        "",
        "",
        "",
    ),
    (
        "TC-AH-011",
        "Strict Rule #3",
        "If information is missing or unclear, respond with 'Insufficient information to determine.'",
        "Verify LLM responds with exact phrase for ambiguous input",
        "Provide ambiguous test data: 'The system has some validation rules.' Ask: 'What are all the validation rules?'",
        "LLM must respond with: 'Insufficient information to determine.'",
        "",
        "",
        "",
    ),
    (
        "TC-AH-012",
        "Strict Rule #3",
        "If information is missing or unclear, respond with 'Insufficient information to determine.'",
        "Verify LLM uses exact phrase for partially-provided info",
        "Provide an API doc showing POST /login with request body 'username'. Ask: 'What is the response format?'",
        "LLM must respond with: 'Insufficient information to determine.' since response format is not provided.",
        "",
        "",
        "",
    ),
    # STRICT RULE 4: Traceable assertions
    (
        "TC-AH-013",
        "Strict Rule #4",
        "Every assertion must be traceable to provided input",
        "Verify LLM can trace each assertion back to a source",
        "Provide a PRD stating 'Login page has email and password fields with a Login button.' Ask: 'Describe the login page.'",
        "LLM response must only mention email field, password field, and Login button. Each element must be traceable.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-014",
        "Strict Rule #4",
        "Every assertion must be traceable to provided input",
        "Verify LLM does NOT make untraceable claims in multi-source scenario",
        "Provide PRD mentioning 'User can view orders' and API docs showing GET /orders. Ask: 'Can users cancel orders?'",
        "LLM must NOT claim users can cancel orders. Only 'view orders' is traceable through both PRD and API docs.",
        "",
        "",
        "",
    ),
    # STRICT RULE 5: Label inferences
    (
        "TC-AH-015",
        "Strict Rule #5",
        "If a detail is inferred, label it explicitly as 'Inference (low confidence)'",
        "Verify LLM labels logical inferences correctly",
        "Provide API docs: POST /login accepts { email, password } and returns { token }. Ask: 'Is the token a JWT?'",
        "If LLM infers it might be a JWT, it MUST label it: 'Inference (low confidence)'. If not inferring, should state insufficient info.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-016",
        "Strict Rule #5",
        "If a detail is inferred, label it explicitly as 'Inference (low confidence)'",
        "Verify LLM labels UI layout inferences",
        "Provide a screenshot with a form and a 'Next' button. Ask: 'What happens when I click Next?'",
        "If LLM infers it navigates to the next page, it MUST label: 'Inference (low confidence)'. The actual behavior is not verifiable from the screenshot.",
        "",
        "",
        "",
    ),
    # STRICT RULE 6: Deterministic output
    (
        "TC-AH-017",
        "Strict Rule #6",
        "Output must be deterministic and repeatable",
        "Verify LLM gives the same answer for the same input across runs",
        "Provide the same PRD describing 'Login with email and password.' Ask twice in separate sessions: 'What fields are on the login page?'",
        "LLM must give the same answer in both runs: should only mention email and password fields each time.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-018",
        "Strict Rule #6",
        "Output must be deterministic and repeatable",
        "Verify LLM does NOT introduce random variations in output",
        "Provide API docs: 'POST /search accepts { query }. Returns { results: [...] }.' Run the same prompt twice.",
        "LLM output must be identical in core facts across both runs. No additional unsourced details should appear.",
        "",
        "",
        "",
    ),
    # PROCESS STEPS
    (
        "TC-AH-019",
        "Process - Step 1",
        "Step 1: Extract verifiable facts from the input",
        "Verify LLM extracts only facts present in the input",
        "Provide: 'The app has 3 user roles: Admin, Editor, Viewer. Admin can delete users.' Ask: 'Summarize user roles.'",
        "LLM must extract exactly: Admin (can delete users), Editor, Viewer. Must NOT add extra permissions for Editor/Viewer.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-020",
        "Process - Step 2",
        "Step 2: List unknown or missing information",
        "Verify LLM identifies and reports missing information",
        "Provide API docs with POST /register { name, email } but no response schema. Ask: 'What does the registration endpoint return?'",
        "LLM must list 'Response schema' as missing/unknown information and NOT fabricate a response body.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-021",
        "Process - Step 3",
        "Step 3: Generate output ONLY from Step 1 facts",
        "Verify LLM output is strictly based on extracted facts",
        "Provide: 'Users table has columns: id, name, email.' Ask: 'What is the database schema for Users?'",
        "LLM must only list id, name, email. Must NOT add created_at, updated_at, or other common columns.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-022",
        "Process - Step 4",
        "Step 4: Perform a self-check for hallucinations or contradictions",
        "Verify LLM self-checks and catches its own potential hallucination",
        "Provide minimal input: 'The login API exists.' Ask: 'Write test cases for login API.'",
        "If LLM generates any output beyond stating insufficient info, it must self-validate. Ideally, it should state 'Insufficient information to determine.'",
        "",
        "",
        "",
    ),
    # OUTPUT FORMAT
    (
        "TC-AH-023",
        "Output Format",
        "Output must follow: Verified Facts / Missing-Unknown / Generated Output / Self-Validation Check",
        "Verify LLM output adheres to the required 4-section format",
        "Provide a PRD describing 'A shopping cart with Add Item and Remove Item features.' Ask: 'What are the cart features?'",
        "LLM output must have these 4 sections: 'Verified Facts:', 'Missing / Unknown Information:', 'Generated Output:', 'Self-Validation Check:'",
        "",
        "",
        "",
    ),
    (
        "TC-AH-024",
        "Output Format",
        "Output must follow: Verified Facts / Missing-Unknown / Generated Output / Self-Validation Check",
        "Verify LLM correctly separates facts from unknowns",
        "Provide test data: 'Login test data: validUser@test.com / Pass123, invalidUser@test.com / wrong' Ask: 'Analyze this test data.'",
        "Verified Facts must list the exact credentials. Missing/Unknown must note what's not specified (e.g., locked accounts, expired passwords).",
        "",
        "",
        "",
    ),
    # EDGE CASES
    (
        "TC-AH-025",
        "Edge Case",
        "If you cannot complete a step, stop and report why.",
        "Verify LLM stops processing when a step cannot be completed",
        "Provide NO input sources and ask: 'Generate test cases for the login feature.'",
        "LLM must stop and report: no input provided, cannot proceed with Step 1 (extracting facts). Must NOT fabricate test cases.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-026",
        "Edge Case",
        "If you cannot complete a step, stop and report why.",
        "Verify LLM handles contradictory input correctly",
        "Provide PRD: 'Users can have multiple addresses.' and API docs: 'POST /users accepts only one address.' Ask: 'How many addresses can a user have?'",
        "LLM must note the contradiction and stop/report ambiguity rather than picking one source over the other.",
        "",
        "",
        "",
    ),
    (
        "TC-AH-027",
        "Edge Case",
        "Combined rules - No invention + No assumption + Traceability",
        "Verify LLM handles complex scenario with multiple rules",
        "Provide: 'A payment page exists with a card number field and a Pay button.' Ask: 'What payment methods are supported?'",
        "LLM must only mention card payment (traceable). Must NOT list UPI, PayPal, net banking. Must NOT assume CVV/expiry fields exist.",
        "",
        "",
        "",
    ),
]

for row_idx, tc in enumerate(test_cases, start=2):
    for col_idx, value in enumerate(tc, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = cell_font
        cell.alignment = wrap_alignment
        cell.border = thin_border
    ws.row_dimensions[row_idx].height = 60

# Add Summary sheet
ws2 = wb.create_sheet("Summary")
ws2.merge_cells("A1:D1")
ws2["A1"] = "Anti-Hallucination Test Cases - Summary"
ws2["A1"].font = Font(name="Calibri", size=14, bold=True, color="2F5496")
ws2["A1"].alignment = Alignment(horizontal="center")

summary_data = [
    ("Module", "Anti-Hallucination Rules for LLM QA Assistant"),
    ("Author", "Pramod Dutta - Principal SDET"),
    ("Total Test Cases", len(test_cases)),
    ("Scope of Knowledge Tests", "3 (TC-AH-001 to TC-AH-003)"),
    ("Strict Rule #1 (No invention) Tests", "3 (TC-AH-004 to TC-AH-006)"),
    ("Strict Rule #2 (No assumptions) Tests", "3 (TC-AH-007 to TC-AH-009)"),
    ("Strict Rule #3 (Insufficient info) Tests", "3 (TC-AH-010 to TC-AH-012)"),
    ("Strict Rule #4 (Traceability) Tests", "2 (TC-AH-013 to TC-AH-014)"),
    ("Strict Rule #5 (Label inferences) Tests", "2 (TC-AH-015 to TC-AH-016)"),
    ("Strict Rule #6 (Deterministic) Tests", "2 (TC-AH-017 to TC-AH-018)"),
    ("Process Steps Tests", "4 (TC-AH-019 to TC-AH-022)"),
    ("Output Format Tests", "2 (TC-AH-023 to TC-AH-024)"),
    ("Edge Case Tests", "3 (TC-AH-025 to TC-AH-027)"),
]

for row_idx, (label, value) in enumerate(summary_data, start=3):
    cell_label = ws2.cell(row=row_idx, column=1, value=label)
    cell_label.font = Font(name="Calibri", size=11, bold=True)
    cell_value = ws2.cell(row=row_idx, column=2, value=value)
    cell_value.font = Font(name="Calibri", size=11)
    if row_idx == 3:
        ws2.merge_cells(f"B{row_idx}:D{row_idx}")
    else:
        ws2.merge_cells(f"B{row_idx}:D{row_idx}")

ws2.column_dimensions["A"].width = 35
ws2.column_dimensions["B"].width = 30
ws2.column_dimensions["C"].width = 15
ws2.column_dimensions["D"].width = 15

output_path = r"C:\Users\kisho\PlayAI4xTester\Test\Anti_Hallucination_Test_Cases.xlsx"
wb.save(output_path)
print(f"Excel file created: {output_path}")
print(f"Total test cases: {len(test_cases)}")
